from collections import defaultdict
from flask import Blueprint, render_template, request, send_file, flash
from flask_login import login_required
from app import db
from app.models import WasteRecord, WasteType, Location
from app.forms import ReportFilterForm
import io
import os
from datetime import date

reports_bp = Blueprint('reports', __name__, url_prefix='/reports')

ORG_NAME    = 'Distributed Environmental Systems'
ORG_ADDRESS = 'Abule Egba, Labak Estate, Lagos, Nigeria'
ORG_PHONE   = '+2348081625419'
ORG_EMAIL   = 'distributedenvironmentalsystems@gmail.com'


def _populate_filter_choices(form):
    form.waste_type_id.choices = [(0, 'All Types')] + \
        [(t.id, t.name) for t in WasteType.query.order_by('name').all()]
    form.location_id.choices = [(0, 'All Locations')] + \
        [(l.id, l.name) for l in Location.query.order_by('name').all()]


def _build_query(form):
    query = db.session.query(
        WasteRecord.date_collected,
        WasteRecord.quantity_kg,
        WasteRecord.vehicle_id,
        WasteRecord.disposal_method,
        WasteRecord.notes,
        WasteType.name.label('waste_type'),
        Location.name.label('location'),
    ).join(WasteType).join(Location)

    if form.date_from.data:
        query = query.filter(WasteRecord.date_collected >= form.date_from.data)
    if form.date_to.data:
        query = query.filter(WasteRecord.date_collected <= form.date_to.data)
    if form.waste_type_id.data and form.waste_type_id.data != 0:
        query = query.filter(WasteRecord.waste_type_id == form.waste_type_id.data)
    if form.location_id.data and form.location_id.data != 0:
        query = query.filter(WasteRecord.location_id == form.location_id.data)

    return query.order_by(WasteRecord.date_collected.desc())


def _build_summary(rows):
    """Build summary dict from raw rows without pandas."""
    if not rows:
        return {}
    total_kg = round(sum(r.quantity_kg for r in rows), 2)

    by_type = defaultdict(float)
    by_location = defaultdict(float)
    by_disposal = defaultdict(float)

    for r in rows:
        by_type[r.waste_type]               += r.quantity_kg
        by_location[r.location]             += r.quantity_kg
        by_disposal[r.disposal_method or 'Unspecified'] += r.quantity_kg

    return {
        'total_kg':      total_kg,
        'total_records': len(rows),
        'by_type':       {k: round(v, 2) for k, v in by_type.items()},
        'by_location':   {k: round(v, 2) for k, v in by_location.items()},
        'by_disposal':   {k: round(v, 2) for k, v in by_disposal.items()},
    }


@reports_bp.route('/', methods=['GET', 'POST'])
@login_required
def index():
    form = ReportFilterForm()
    _populate_filter_choices(form)
    records = []
    summary = {}

    if form.validate_on_submit() or request.method == 'GET':
        rows    = _build_query(form).all()
        records = rows
        summary = _build_summary(rows)

    return render_template('reports/index.html', title='Reports',
                           form=form, records=records, summary=summary)


@reports_bp.route('/export/excel', methods=['POST'])
@login_required
def export_excel():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    form = ReportFilterForm()
    _populate_filter_choices(form)
    form.validate()

    rows = _build_query(form).all()
    if not rows:
        flash('No records to export.', 'warning')
        return render_template('reports/index.html', form=form, records=[], summary={})

    wb = Workbook()
    ws = wb.active
    ws.title = 'Waste Records'

    # Header row styling
    headers = ['Date', 'Waste Type', 'Location', 'Vehicle ID',
               'Disposal Method', 'Quantity (kg)', 'Notes']
    green_fill = PatternFill('solid', fgColor='198754')
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = Font(bold=True, color='FFFFFF')
        cell.fill      = green_fill
        cell.alignment = Alignment(horizontal='center')

    # Data rows
    for row_idx, r in enumerate(rows, 2):
        ws.append([
            str(r.date_collected), r.waste_type, r.location,
            r.vehicle_id or '', r.disposal_method or '',
            r.quantity_kg, r.notes or ''
        ])

    # Summary sheet
    ws2 = wb.create_sheet('Summary')
    summary = _build_summary(rows)
    ws2.append(['Metric', 'Value'])
    ws2.append(['Total Records', summary['total_records']])
    ws2.append(['Total Weight (kg)', summary['total_kg']])
    ws2.append([])
    ws2.append(['By Waste Type', ''])
    for k, v in summary['by_type'].items():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(['By Disposal Method', ''])
    for k, v in summary['by_disposal'].items():
        ws2.append([k, v])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f'waste_report_{date.today()}.xlsx'
    return send_file(output, download_name=filename, as_attachment=True,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@reports_bp.route('/export/pdf', methods=['POST'])
@login_required
def export_pdf():
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import (SimpleDocTemplate, Table, TableStyle,
                                    Paragraph, Spacer, HRFlowable, Image)
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_CENTER

    form = ReportFilterForm()
    _populate_filter_choices(form)
    form.validate()

    rows = _build_query(form).all()
    if not rows:
        flash('No records to export.', 'warning')
        return render_template('reports/index.html', form=form, records=[], summary={})

    output  = io.BytesIO()
    doc     = SimpleDocTemplate(output, pagesize=landscape(A4),
                                 leftMargin=1.5*cm, rightMargin=1.5*cm,
                                 topMargin=1.5*cm, bottomMargin=1.5*cm)
    GREEN     = colors.HexColor('#198754')
    LIGHT_ROW = colors.HexColor('#f0faf4')
    GREY_LINE = colors.HexColor('#dee2e6')

    org_style   = ParagraphStyle('org',   fontSize=16, textColor=GREEN,
                                  fontName='Helvetica-Bold', spaceAfter=2)
    sub_style   = ParagraphStyle('sub',   fontSize=9,
                                  textColor=colors.HexColor('#6c757d'), spaceAfter=1)
    title_style = ParagraphStyle('title', fontSize=12, fontName='Helvetica-Bold',
                                  spaceBefore=6, spaceAfter=4)

    elements = []

    # Logo + org header
    logo_path = os.path.normpath(
        os.path.join(os.path.dirname(__file__), '..', 'static', 'images', 'logo.png'))
    if os.path.exists(logo_path):
        logo = Image(logo_path, width=2*cm, height=2*cm)
        hdata = [[logo, [Paragraph(ORG_NAME, org_style),
                         Paragraph(ORG_ADDRESS, sub_style),
                         Paragraph(f'Tel: {ORG_PHONE}', sub_style),
                         Paragraph(f'Email: {ORG_EMAIL}', sub_style)]]]
        htable = Table(hdata, colWidths=[2.5*cm, None])
    else:
        hdata = [[[Paragraph(ORG_NAME, org_style),
                   Paragraph(ORG_ADDRESS, sub_style),
                   Paragraph(f'Tel: {ORG_PHONE}', sub_style),
                   Paragraph(f'Email: {ORG_EMAIL}', sub_style)]]]
        htable = Table(hdata)

    htable.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
    ]))
    elements.append(htable)
    elements.append(HRFlowable(width='100%', thickness=1.5, color=GREEN, spaceAfter=6))

    date_range = ''
    if form.date_from.data or form.date_to.data:
        d_from = str(form.date_from.data) if form.date_from.data else 'All time'
        d_to   = str(form.date_to.data)   if form.date_to.data   else 'Present'
        date_range = f'  |  Period: {d_from} to {d_to}'

    elements.append(Paragraph(f'Waste Management Report{date_range}', title_style))
    elements.append(Paragraph(
        f'Generated: {date.today()}   |   Total records: {len(rows)}',
        ParagraphStyle('meta', fontSize=8, textColor=colors.HexColor('#6c757d'), spaceAfter=8)))

    summary   = _build_summary(rows)
    sum_items = [f'Total: {summary["total_kg"]} kg'] + \
                [f'{k}: {v} kg' for k, v in summary['by_type'].items()]
    elements.append(Paragraph('  |  '.join(sum_items),
                               ParagraphStyle('summ', fontSize=8, textColor=GREEN,
                                              fontName='Helvetica-Bold', spaceAfter=8)))
    elements.append(HRFlowable(width='100%', thickness=0.5, color=GREY_LINE, spaceAfter=6))

    tdata = [['Date', 'Waste Type', 'Location', 'Vehicle ID', 'Disposal', 'Qty (kg)', 'Notes']]
    for r in rows:
        tdata.append([str(r.date_collected), r.waste_type, r.location,
                      r.vehicle_id or '—', r.disposal_method or '—',
                      f'{r.quantity_kg:.2f}', (r.notes or '')[:40]])

    table = Table(tdata, colWidths=[2.5*cm,3*cm,4*cm,2.5*cm,2.8*cm,2*cm,None], repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND',    (0,0), (-1,0), GREEN),
        ('TEXTCOLOR',     (0,0), (-1,0), colors.white),
        ('FONTNAME',      (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE',      (0,0), (-1,0), 9),
        ('ROWBACKGROUNDS',(0,1), (-1,-1), [colors.white, LIGHT_ROW]),
        ('FONTSIZE',      (0,1), (-1,-1), 8),
        ('GRID',          (0,0), (-1,-1), 0.4, GREY_LINE),
        ('VALIGN',        (0,0), (-1,-1), 'MIDDLE'),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 10))
    elements.append(HRFlowable(width='100%', thickness=0.5, color=GREY_LINE))
    elements.append(Paragraph(
        f'{ORG_NAME}  |  {ORG_ADDRESS}  |  {ORG_PHONE}  |  {ORG_EMAIL}',
        ParagraphStyle('footer', fontSize=7, textColor=colors.HexColor('#6c757d'),
                       alignment=TA_CENTER, spaceBefore=4)))

    doc.build(elements)
    output.seek(0)
    return send_file(io.BytesIO(output.getvalue()),
                     download_name=f'waste_report_{date.today()}.pdf',
                     as_attachment=True, mimetype='application/pdf')
